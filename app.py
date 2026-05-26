import asyncio
import csv
import io
import os
import re
from datetime import datetime, date, timedelta

import hashlib
from flask import Flask, render_template, request, redirect, url_for, flash, Response, jsonify, session, send_file
from flask_sqlalchemy import SQLAlchemy
from flask_login import LoginManager, login_user, logout_user, login_required, UserMixin, current_user
from apscheduler.schedulers.background import BackgroundScheduler
from dotenv import load_dotenv

load_dotenv()

app = Flask(__name__)
app.config["SECRET_KEY"] = os.environ.get("SECRET_KEY", "dev-secret-change-me")
_db_path = os.environ.get("DATABASE_URL", "sqlite:////data/ventes.db")
if _db_path.startswith("postgres://"):
    _db_path = _db_path.replace("postgres://", "postgresql://", 1)
elif "sqlite" in _db_path:
    m = re.match(r'sqlite:////(.+)', _db_path)
    if m:
        try:
            os.makedirs(os.path.dirname("/" + m.group(1)), exist_ok=True)
        except Exception:
            _db_path = "sqlite:////tmp/ventes.db"
            os.makedirs("/tmp", exist_ok=True)
app.config["SQLALCHEMY_DATABASE_URI"] = _db_path
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False

db = SQLAlchemy(app)

login_manager = LoginManager(app)
login_manager.login_view = "login"
login_manager.login_message = "Connecte-toi pour accéder à tes ventes."
login_manager.login_message_category = "warning"

APP_PASSWORD_HASH = hashlib.sha256(
    os.environ.get("APP_PASSWORD", "orange2026").encode()
).hexdigest()


class FakeUser(UserMixin):
    id = "1"


THE_USER = FakeUser()


@login_manager.user_loader
def load_user(user_id):
    return THE_USER if user_id == "1" else None


@app.route("/login", methods=["GET", "POST"])
def login():
    if current_user.is_authenticated:
        return redirect(url_for("dashboard"))
    if request.method == "POST":
        pwd = request.form.get("password", "")
        if hashlib.sha256(pwd.encode()).hexdigest() == APP_PASSWORD_HASH:
            login_user(THE_USER, remember=True)
            return redirect(request.args.get("next") or url_for("dashboard"))
        flash("Mot de passe incorrect.", "danger")
    return render_template("login.html")


@app.route("/logout")
@login_required
def logout():
    logout_user()
    return redirect(url_for("login"))


STATUTS = {
    "en_attente": "En attente",
    "confirme": "Confirmé",
    "installe": "Installé",
    "no_show": "No-show",
    "annule": "Annulé",
}

PRODUITS = [
    "En option",
    "Livebox Fibre",
    "Livebox Up",
    "Livebox Max",
    "Série Spécial Lite Fibre",
]


class Vente(db.Model):
    __tablename__ = "ventes"

    id = db.Column(db.Integer, primary_key=True)
    prenom = db.Column(db.String(100), nullable=False)
    nom = db.Column(db.String(100), nullable=False)
    telephone = db.Column(db.String(20), nullable=False)
    adresse = db.Column(db.String(200), nullable=False)
    produit = db.Column(db.String(100), nullable=False)
    date_rdv = db.Column(db.DateTime, nullable=False)
    date_signature = db.Column(db.Date, nullable=False, default=date.today)
    statut = db.Column(db.String(50), nullable=False, default="en_attente")
    reference = db.Column(db.String(100), nullable=True)
    sms_envoye = db.Column(db.Boolean, default=False)
    notes = db.Column(db.Text, nullable=True)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)

    @property
    def statut_label(self):
        return STATUTS.get(self.statut, self.statut)

    @property
    def badge_class(self):
        return {
            "en_attente": "bg-warning text-dark",
            "confirme": "bg-info text-dark",
            "installe": "bg-success",
            "no_show": "bg-danger",
            "annule": "bg-secondary",
        }.get(self.statut, "bg-secondary")


RESULTATS_PORTE = {
    "absent":  {"label": "ABSENT",  "emoji": "🔘", "color": "#6c757d"},
    "refus":   {"label": "REFUS",   "emoji": "❌", "color": "#dc3545"},
    "cause":   {"label": "CAUSÉ",   "emoji": "💬", "color": "#fd7e14"},
    "entre":   {"label": "ENTRÉ",   "emoji": "🏠", "color": "#0d6efd"},
    "signe":   {"label": "SIGNÉ",   "emoji": "✅", "color": "#198754"},
}


class SessionProspection(db.Model):
    __tablename__ = "sessions_prospection"

    id = db.Column(db.Integer, primary_key=True)
    nom = db.Column(db.String(200), nullable=False)
    date = db.Column(db.Date, nullable=False, default=date.today)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    portes = db.relationship(
        "Porte", backref="session", lazy=True,
        cascade="all, delete-orphan", order_by="Porte.id",
    )

    @property
    def total(self):
        return len(self.portes)

    @property
    def nb_ouvertes(self):
        return sum(1 for p in self.portes if p.resultat in ("refus", "cause", "entre", "signe"))

    @property
    def nb_causes(self):
        return sum(1 for p in self.portes if p.resultat in ("cause", "entre", "signe"))

    @property
    def nb_entrees(self):
        return sum(1 for p in self.portes if p.resultat in ("entre", "signe"))

    @property
    def nb_signes(self):
        return sum(1 for p in self.portes if p.resultat == "signe")

    def _pct(self, n):
        return round(n / self.total * 100) if self.total else 0

    @property
    def taux_ouverture(self):
        return self._pct(self.nb_ouvertes)

    @property
    def taux_entree(self):
        return self._pct(self.nb_entrees)

    @property
    def taux_cause(self):
        return self._pct(self.nb_causes)

    @property
    def taux_signature(self):
        return self._pct(self.nb_signes)

    def to_stats_dict(self):
        return {
            "total":          self.total,
            "nb_ouvertes":    self.nb_ouvertes,
            "nb_causes":      self.nb_causes,
            "nb_entrees":     self.nb_entrees,
            "nb_signes":      self.nb_signes,
            "taux_ouverture": self.taux_ouverture,
            "taux_cause":     self.taux_cause,
            "taux_entree":    self.taux_entree,
            "taux_signature": self.taux_signature,
        }


class Porte(db.Model):
    __tablename__ = "portes"

    id = db.Column(db.Integer, primary_key=True)
    session_id = db.Column(
        db.Integer, db.ForeignKey("sessions_prospection.id"), nullable=False
    )
    resultat = db.Column(db.String(20), nullable=False)
    adresse_id = db.Column(db.Integer, db.ForeignKey("adresses_importees.id"), nullable=True)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)


class AdresseImportee(db.Model):
    __tablename__ = "adresses_importees"

    id = db.Column(db.Integer, primary_key=True)
    rue = db.Column(db.String(200), nullable=False, index=True)
    numero = db.Column(db.String(20), nullable=False)
    complement = db.Column(db.String(100), nullable=True)
    ville = db.Column(db.String(100), nullable=True, index=True)
    imported_at = db.Column(db.DateTime, default=datetime.utcnow)

    @property
    def adresse_complete(self):
        parts = [self.numero, self.rue]
        if self.complement:
            parts.append(self.complement)
        return " ".join(parts)


MOMENTS_RAPPEL = {
    "matin":  "🌅 Matin",
    "pause":  "☀️ Pause",
    "soir":   "🌙 Soir",
}


class Rappel(db.Model):
    __tablename__ = "rappels"

    id = db.Column(db.Integer, primary_key=True)
    nom = db.Column(db.String(100), nullable=False)
    telephone = db.Column(db.String(20), nullable=True)
    motif = db.Column(db.String(200), nullable=False)
    moment = db.Column(db.String(20), nullable=False, default="pause")
    done = db.Column(db.Boolean, default=False)
    vente_id = db.Column(db.Integer, db.ForeignKey("ventes.id"), nullable=True)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)


class SalaireMensuel(db.Model):
    __tablename__ = "salaires_mensuels"

    id = db.Column(db.Integer, primary_key=True)
    annee = db.Column(db.Integer, nullable=False)
    mois = db.Column(db.Integer, nullable=False)
    montant_net = db.Column(db.Float, nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    __table_args__ = (db.UniqueConstraint("annee", "mois", name="uq_salaire_mois"),)


def envoyer_sms(telephone: str, message: str) -> bool:
    sid = os.environ.get("TWILIO_ACCOUNT_SID")
    token = os.environ.get("TWILIO_AUTH_TOKEN")
    from_number = os.environ.get("TWILIO_PHONE_NUMBER")
    if not all([sid, token, from_number]):
        app.logger.warning("Twilio non configuré — SMS non envoyé.")
        return False
    try:
        from twilio.rest import Client
        client = Client(sid, token)
        client.messages.create(body=message, from_=from_number, to=telephone)
        return True
    except Exception as exc:
        app.logger.error("Erreur Twilio : %s", exc)
        return False


def construire_message_rappel(vente: Vente) -> str:
    date_str = vente.date_rdv.strftime("%d/%m/%Y")
    heure_str = vente.date_rdv.strftime("%H:%M")
    return (
        f"Bonjour {vente.prenom},\n"
        f"Rappel Orange : votre rendez-vous d'installation ({vente.produit}) "
        f"est prévu demain {date_str} à {heure_str}.\n"
        f"Pour modifier, appelez le 3900.\n"
        f"À bientôt !"
    )


def envoyer_sms_rappels_clients():
    with app.app_context():
        mon_tel = os.environ.get("MON_TELEPHONE", "")
        if not mon_tel:
            return
        rappels = Rappel.query.filter_by(done=False).order_by(Rappel.created_at).all()
        if not rappels:
            return
        lignes = [f"📋 {len(rappels)} rappel(s) client en attente :"]
        for r in rappels:
            tel_part = f" ({r.telephone})" if r.telephone else ""
            emoji = {"matin": "🌅", "pause": "☀️", "soir": "🌙"}.get(r.moment, "")
            lignes.append(f"{emoji} {r.nom}{tel_part} : {r.motif}")
        envoyer_sms(mon_tel, "\n".join(lignes))


def envoyer_rappels_du_jour():
    with app.app_context():
        demain = date.today() + timedelta(days=1)
        ventes = Vente.query.filter(
            db.func.date(Vente.date_rdv) == demain,
            Vente.sms_envoye == False,
            Vente.statut.in_(["en_attente", "confirme"]),
        ).all()
        for vente in ventes:
            message = construire_message_rappel(vente)
            if envoyer_sms(vente.telephone, message):
                vente.sms_envoye = True
                db.session.commit()
                app.logger.info("SMS rappel envoyé → %s %s", vente.prenom, vente.nom)


@app.context_processor
def inject_nav_badges():
    if current_user.is_authenticated:
        nb_rappels = Rappel.query.filter_by(done=False).count()
        nb_repasser = (
            Porte.query
            .filter_by(resultat="absent")
            .filter(Porte.adresse_id.isnot(None))
            .count()
        )
        return dict(nb_rappels=nb_rappels, nb_repasser=nb_repasser)
    return dict(nb_rappels=0, nb_repasser=0)


@app.route("/recap")
@app.route("/recap/<int:offset_semaines>")
@login_required
def recap_semaine(offset_semaines=0):
    aujourd_hui = date.today()
    lundi = aujourd_hui - timedelta(days=aujourd_hui.weekday()) + timedelta(weeks=offset_semaines)
    dimanche = lundi + timedelta(days=6)
    ventes_semaine = Vente.query.filter(
        Vente.date_signature >= lundi,
        Vente.date_signature <= dimanche,
    ).order_by(Vente.date_signature.asc(), Vente.created_at.asc()).all()
    jours = []
    noms_jours = ["Lundi", "Mardi", "Mercredi", "Jeudi", "Vendredi", "Samedi", "Dimanche"]
    for i in range(7):
        jour = lundi + timedelta(days=i)
        ventes_jour = [v for v in ventes_semaine if v.date_signature == jour]
        jours.append({
            "nom": noms_jours[i],
            "date": jour,
            "ventes": ventes_jour,
            "is_today": jour == aujourd_hui,
        })
    par_produit = {}
    for v in ventes_semaine:
        par_produit[v.produit] = par_produit.get(v.produit, 0) + 1
    return render_template(
        "recap.html",
        jours=jours,
        lundi=lundi,
        dimanche=dimanche,
        total_semaine=len(ventes_semaine),
        par_produit=par_produit,
        offset=offset_semaines,
        aujourd_hui=aujourd_hui,
    )


def get_niveau_info(n):
    if n >= 100:
        return dict(nom="Diamant", emoji="💎", color="#0dcaf0", progress=100,
                    manquant=0, prochain=None, prochain_emoji=None)
    elif n >= 50:
        return dict(nom="Or", emoji="🥇", color="#ffc107",
                    progress=round((n - 50) / 50 * 100), manquant=100 - n,
                    prochain="Diamant", prochain_emoji="💎")
    elif n >= 20:
        return dict(nom="Argent", emoji="🥈", color="#adb5bd",
                    progress=round((n - 20) / 30 * 100), manquant=50 - n,
                    prochain="Or", prochain_emoji="🥇")
    else:
        return dict(nom="Bronze", emoji="🥉", color="#b87333",
                    progress=round(n / 20 * 100) if n else 0, manquant=20 - n,
                    prochain="Argent", prochain_emoji="🥈")


@app.route("/")
@login_required
def dashboard():
    from collections import defaultdict
    aujourd_hui = date.today()
    demain = aujourd_hui + timedelta(days=1)
    total = Vente.query.count()
    rdv_demain = Vente.query.filter(
        db.func.date(Vente.date_rdv) == demain,
        Vente.statut.in_(["en_attente", "confirme"]),
    ).count()
    no_shows = Vente.query.filter_by(statut="no_show").count()
    installes = Vente.query.filter_by(statut="installe").count()
    en_attente = Vente.query.filter_by(statut="en_attente").count()
    taux_no_show = round((no_shows / total * 100) if total else 0, 1)
    ventes_recentes = Vente.query.order_by(Vente.created_at.desc()).limit(8).all()
    rdv_demain_liste = Vente.query.filter(
        db.func.date(Vente.date_rdv) == demain,
        Vente.statut.in_(["en_attente", "confirme"]),
    ).order_by(Vente.date_rdv).all()
    objectif = int(os.environ.get("OBJECTIF_MENSUEL", "20"))
    debut_mois = aujourd_hui.replace(day=1)
    ventes_ce_mois = Vente.query.filter(Vente.date_signature >= debut_mois).count()
    progress_objectif = min(100, round(ventes_ce_mois / objectif * 100)) if objectif else 0
    niveau = get_niveau_info(total)
    ventes_par_jour = defaultdict(int)
    for (ds,) in db.session.query(Vente.date_signature).all():
        ventes_par_jour[ds] += 1
    record_jour = max(ventes_par_jour.values()) if ventes_par_jour else 0
    ventes_par_semaine = defaultdict(int)
    for d, nb in ventes_par_jour.items():
        lundi = d - timedelta(days=d.weekday())
        ventes_par_semaine[lundi] += nb
    record_semaine = max(ventes_par_semaine.values()) if ventes_par_semaine else 0
    lundi_cette_sem = aujourd_hui - timedelta(days=aujourd_hui.weekday())
    lundi_sem_prec  = lundi_cette_sem - timedelta(weeks=1)
    dim_sem_prec    = lundi_cette_sem - timedelta(days=1)
    ventes_cette_sem = Vente.query.filter(Vente.date_signature >= lundi_cette_sem).count()
    ventes_sem_prec = Vente.query.filter(
        Vente.date_signature >= lundi_sem_prec,
        Vente.date_signature <= dim_sem_prec,
    ).count()
    diff_semaine = ventes_cette_sem - ventes_sem_prec
    return render_template(
        "dashboard.html",
        total=total, rdv_demain=rdv_demain, no_shows=no_shows,
        installes=installes, en_attente=en_attente, taux_no_show=taux_no_show,
        ventes_recentes=ventes_recentes, rdv_demain_liste=rdv_demain_liste,
        aujourd_hui=aujourd_hui, objectif=objectif, ventes_ce_mois=ventes_ce_mois,
        progress_objectif=progress_objectif, niveau=niveau, record_jour=record_jour,
        record_semaine=record_semaine, ventes_cette_sem=ventes_cette_sem,
        ventes_sem_prec=ventes_sem_prec, diff_semaine=diff_semaine,
    )


@app.route("/ventes")
@login_required
def liste_ventes():
    statut_filtre = request.args.get("statut", "")
    produit_filtre = request.args.get("produit", "")
    search = request.args.get("q", "").strip()
    query = Vente.query
    if statut_filtre:
        query = query.filter_by(statut=statut_filtre)
    if produit_filtre:
        query = query.filter_by(produit=produit_filtre)
    if search:
        like = f"%{search}%"
        query = query.filter(
            db.or_(
                Vente.prenom.ilike(like), Vente.nom.ilike(like),
                Vente.telephone.ilike(like), Vente.adresse.ilike(like),
            )
        )
    ventes = query.order_by(Vente.date_rdv.desc()).all()
    if request.args.get("export") == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["ID", "Prénom", "Nom", "Téléphone", "Adresse",
            "Produit", "Réf. commande", "Date RDV", "Date signature", "Statut", "SMS envoyé", "Notes"])
        for v in ventes:
            writer.writerow([v.id, v.prenom, v.nom, v.telephone, v.adresse, v.produit,
                v.reference or "", v.date_rdv.strftime("%d/%m/%Y %H:%M"),
                v.date_signature.strftime("%d/%m/%Y"), v.statut_label,
                "Oui" if v.sms_envoye else "Non", v.notes or ""])
        output.seek(0)
        return Response("﻿" + output.getvalue(), mimetype="text/csv",
            headers={"Content-Disposition": "attachment; filename=ventes_orange.csv"})
    return render_template("ventes.html", ventes=ventes, statuts=STATUTS, produits=PRODUITS,
        statut_filtre=statut_filtre, produit_filtre=produit_filtre, search=search)


@app.route("/ajouter", methods=["GET", "POST"])
@login_required
def ajouter_vente():
    if request.method == "POST":
        try:
            date_rdv = datetime.strptime(request.form["date_rdv"], "%Y-%m-%dT%H:%M")
            date_signature_str = request.form.get("date_signature")
            date_signature = (
                datetime.strptime(date_signature_str, "%Y-%m-%d").date()
                if date_signature_str else date.today()
            )
            vente = Vente(
                prenom=request.form["prenom"].strip().capitalize(),
                nom=request.form["nom"].strip().upper(),
                telephone=request.form["telephone"].strip(),
                adresse=request.form["adresse"].strip(),
                produit=request.form["produit"],
                reference=request.form.get("reference", "").strip() or None,
                date_rdv=date_rdv, date_signature=date_signature,
                statut=request.form.get("statut", "en_attente"),
                notes=request.form.get("notes", "").strip() or None,
            )
            db.session.add(vente)
            db.session.commit()
            flash(f"Vente ajoutée : {vente.prenom} {vente.nom}", "success")
            return redirect(url_for("dashboard"))
        except Exception as exc:
            flash(f"Erreur : {exc}", "danger")
    return render_template("formulaire.html", vente=None, produits=PRODUITS,
        statuts=STATUTS, titre="Nouvelle vente")


@app.route("/modifier/<int:vente_id>", methods=["GET", "POST"])
@login_required
def modifier_vente(vente_id):
    vente = Vente.query.get_or_404(vente_id)
    if request.method == "POST":
        try:
            vente.prenom = request.form["prenom"].strip().capitalize()
            vente.nom = request.form["nom"].strip().upper()
            vente.telephone = request.form["telephone"].strip()
            vente.adresse = request.form["adresse"].strip()
            vente.produit = request.form["produit"]
            vente.reference = request.form.get("reference", "").strip() or None
            vente.date_rdv = datetime.strptime(request.form["date_rdv"], "%Y-%m-%dT%H:%M")
            date_sig_str = request.form.get("date_signature")
            if date_sig_str:
                vente.date_signature = datetime.strptime(date_sig_str, "%Y-%m-%d").date()
            ancien_statut = vente.statut
            vente.statut = request.form.get("statut", vente.statut)
            vente.notes = request.form.get("notes", "").strip() or None
            if vente.statut != ancien_statut and vente.statut in ("en_attente", "confirme"):
                vente.sms_envoye = False
            db.session.commit()
            flash("Vente mise à jour.", "success")
            return redirect(url_for("liste_ventes"))
        except Exception as exc:
            flash(f"Erreur : {exc}", "danger")
    return render_template("formulaire.html", vente=vente, produits=PRODUITS,
        statuts=STATUTS, titre="Modifier la vente")


@app.route("/supprimer/<int:vente_id>", methods=["POST"])
@login_required
def supprimer_vente(vente_id):
    vente = Vente.query.get_or_404(vente_id)
    db.session.delete(vente)
    db.session.commit()
    flash("Vente supprimée.", "info")
    return redirect(url_for("liste_ventes"))


@app.route("/statut/<int:vente_id>/<statut>", methods=["POST"])
@login_required
def changer_statut(vente_id, statut):
    if statut not in STATUTS:
        flash("Statut invalide.", "danger")
        return redirect(url_for("liste_ventes"))
    vente = Vente.query.get_or_404(vente_id)
    vente.statut = statut
    db.session.commit()
    flash(f"{vente.prenom} {vente.nom} → {STATUTS[statut]}", "success")
    return redirect(request.referrer or url_for("liste_ventes"))


@app.route("/stats")
@login_required
def stats():
    from collections import defaultdict
    from sqlalchemy import func as _func
    aujourd_hui = date.today()
    total_ventes = Vente.query.count()
    installes   = Vente.query.filter_by(statut="installe").count()
    no_shows    = Vente.query.filter_by(statut="no_show").count()
    annules     = Vente.query.filter_by(statut="annule").count()
    taux_installation = round(installes / total_ventes * 100) if total_ventes else 0
    taux_no_show      = round(no_shows  / total_ventes * 100) if total_ventes else 0
    par_produit = (
        db.session.query(Vente.produit, _func.count(Vente.id))
        .group_by(Vente.produit).order_by(_func.count(Vente.id).desc()).all()
    )
    JOURS = ["Lundi", "Mardi", "Mercredi", "Jeudi", "Vendredi", "Samedi", "Dimanche"]
    par_jour = defaultdict(int)
    for (ds,) in db.session.query(Vente.date_signature).all():
        par_jour[ds.weekday()] += 1
    par_jour_data = [(JOURS[i], par_jour[i]) for i in range(7)]
    meilleur_jour = JOURS[max(range(7), key=lambda i: par_jour[i])] if total_ventes else "—"
    semaines = []
    for i in range(4):
        lundi    = aujourd_hui - timedelta(days=aujourd_hui.weekday()) - timedelta(weeks=i)
        dimanche = lundi + timedelta(days=6)
        nb = Vente.query.filter(
            Vente.date_signature >= lundi, Vente.date_signature <= dimanche,
        ).count()
        semaines.append({"label": "Cette sem." if i == 0 else f"S-{i}", "nb": nb})
    semaines.reverse()
    debut_mois = aujourd_hui.replace(day=1)
    if debut_mois.month == 1:
        debut_mois_dernier = debut_mois.replace(year=debut_mois.year - 1, month=12)
    else:
        debut_mois_dernier = debut_mois.replace(month=debut_mois.month - 1)
    ventes_ce_mois      = Vente.query.filter(Vente.date_signature >= debut_mois).count()
    ventes_mois_dernier = Vente.query.filter(
        Vente.date_signature >= debut_mois_dernier, Vente.date_signature < debut_mois,
    ).count()
    total_portes      = Porte.query.count()
    nb_absent         = Porte.query.filter_by(resultat="absent").count()
    nb_cause          = Porte.query.filter_by(resultat="cause").count()
    nb_entre          = Porte.query.filter_by(resultat="entre").count()
    nb_signe_terrain  = Porte.query.filter_by(resultat="signe").count()
    portes_ouvert      = total_portes - nb_absent
    portes_cause_total = nb_cause + nb_entre + nb_signe_terrain
    taux_ouverture    = round(portes_ouvert      / total_portes * 100) if total_portes else 0
    taux_cause_p      = round(portes_cause_total / total_portes * 100) if total_portes else 0
    taux_entre_p      = round((nb_entre + nb_signe_terrain) / total_portes * 100) if total_portes else 0
    taux_signe_p      = round(nb_signe_terrain   / total_portes * 100) if total_portes else 0
    ratio_portes_vente = round(total_portes / total_ventes) if total_ventes and total_portes else None
    from sqlalchemy import func as _func2
    premiere_vente_date = db.session.query(_func2.min(Vente.date_signature)).scalar()
    nb_jours_actifs = db.session.query(
        _func2.count(_func2.distinct(Vente.date_signature))
    ).scalar() or 0
    if premiere_vente_date and total_ventes:
        nb_jours_calendrier = (aujourd_hui - premiere_vente_date).days + 1
        nb_semaines_calendrier = max(1, round(nb_jours_calendrier / 7))
        nb_mois_calendrier = max(1, round(nb_jours_calendrier / 30))
        moyenne_par_jour    = round(total_ventes / nb_jours_actifs, 1) if nb_jours_actifs else 0
        moyenne_par_semaine = round(total_ventes / nb_semaines_calendrier, 1)
        moyenne_par_mois    = round(total_ventes / nb_mois_calendrier, 1)
    else:
        moyenne_par_jour = moyenne_par_semaine = moyenne_par_mois = 0
        nb_jours_actifs = 0
    MOIS_FR = ["", "Janvier", "Février", "Mars", "Avril", "Mai", "Juin",
               "Juillet", "Août", "Septembre", "Octobre", "Novembre", "Décembre"]
    par_mois_dict = defaultdict(lambda: {"total": 0, "installes": 0, "no_shows": 0, "annules": 0})
    for v in Vente.query.all():
        key = (v.date_signature.year, v.date_signature.month)
        par_mois_dict[key]["total"] += 1
        if v.statut == "installe":
            par_mois_dict[key]["installes"] += 1
        elif v.statut == "no_show":
            par_mois_dict[key]["no_shows"] += 1
        elif v.statut == "annule":
            par_mois_dict[key]["annules"] += 1
    par_mois = []
    salaires = {(s.annee, s.mois): s.montant_net for s in SalaireMensuel.query.all()}
    for (year, month), data in sorted(par_mois_dict.items(), reverse=True):
        t = data["total"]
        taux = round(data["installes"] / t * 100) if t else 0
        par_mois.append({
            "label": f"{MOIS_FR[month]} {year}",
            "annee": year,
            "mois": month,
            "total": t,
            "installes": data["installes"],
            "no_shows": data["no_shows"],
            "annules": data["annules"],
            "taux": taux,
            "salaire_net": salaires.get((year, month)),
        })
    return render_template("stats.html",
        aujourd_hui=aujourd_hui, total_ventes=total_ventes, installes=installes,
        no_shows=no_shows, annules=annules, taux_installation=taux_installation,
        taux_no_show=taux_no_show, par_produit=par_produit, par_jour_data=par_jour_data,
        meilleur_jour=meilleur_jour, ventes_ce_mois=ventes_ce_mois,
        ventes_mois_dernier=ventes_mois_dernier, semaines=semaines,
        total_portes=total_portes, portes_ouvert=portes_ouvert,
        portes_cause_total=portes_cause_total, nb_entre=nb_entre,
        nb_signe_terrain=nb_signe_terrain, taux_ouverture=taux_ouverture,
        taux_cause_p=taux_cause_p, taux_entre_p=taux_entre_p, taux_signe_p=taux_signe_p,
        ratio_portes_vente=ratio_portes_vente, moyenne_par_jour=moyenne_par_jour,
        moyenne_par_semaine=moyenne_par_semaine, moyenne_par_mois=moyenne_par_mois,
        nb_jours_actifs=nb_jours_actifs, par_mois=par_mois)


@app.route("/carte")
@login_required
def carte():
    ventes = Vente.query.order_by(Vente.date_rdv.desc()).all()
    ventes_data = [
        {"id": v.id, "nom": f"{v.prenom} {v.nom}", "adresse": v.adresse,
         "produit": v.produit, "statut": v.statut, "date_rdv": v.date_rdv.strftime("%d/%m/%Y")}
        for v in ventes if v.adresse
    ]
    return render_template("carte.html", ventes=ventes_data)


@app.route("/envoyer-sms/<int:vente_id>", methods=["POST"])
@login_required
def envoyer_sms_manuel(vente_id):
    vente = Vente.query.get_or_404(vente_id)
    message = construire_message_rappel(vente)
    if envoyer_sms(vente.telephone, message):
        vente.sms_envoye = True
        db.session.commit()
        flash(f"SMS envoyé à {vente.prenom} {vente.nom} ({vente.telephone}).", "success")
    else:
        flash("Impossible d'envoyer le SMS. Vérifiez la config Twilio dans .env.", "warning")
    return redirect(request.referrer or url_for("liste_ventes"))


@app.route("/lancer-rappels", methods=["POST"])
@login_required
def lancer_rappels():
    envoyer_rappels_du_jour()
    flash("Rappels SMS du lendemain traités.", "info")
    return redirect(url_for("dashboard"))


_PROMPT_SCAN = (
    "Tu es un assistant pour un commercial Orange en porte-à-porte.\n"
    "Extrait les informations client de ce document "
    "(bon de commande, contrat, écran CRM Orange, suivi de commande, fiche client...).\n\n"
    "Retourne UNIQUEMENT un objet JSON valide avec ces champs (null si non trouvé) :\n"
    "{\n"
    '  "prenom": "...",\n  "nom": "...",\n  "telephone": "...",\n'
    '  "adresse": "...",\n  "produit": "...",\n  "reference": "...",\n'
    '  "date_rdv": "YYYY-MM-DDTHH:MM",\n  "statut": "..."\n'
    "}\n\n"
    "Règles :\n"
    "- nom/prenom : sur l'écran CRM Orange le nom complet est souvent dans la section 'client' "
    "(ex: 'Blanchet Chantale' → prenom='Chantale', nom='Blanchet').\n"
    "- telephone : prendre le numéro mobile (07/06) en priorité. "
    "Convertir '07 83 14 00 32' → '+33783140032'.\n"
    "- produit : choisir le plus proche parmi : "
    "'En option', 'Livebox Fibre', 'Livebox Up', 'Livebox Max', 'Série Spécial Lite Fibre'. "
    "'Livebox Classic Fibre', 'Livebox Fibre +' → 'Livebox Fibre'. "
    "'Livebox Up Fibre' → 'Livebox Up'.\n"
    "- reference : prendre EN PRIORITÉ 'référence interne'. Sinon 'référence commande'. "
    "Ignorer 'code d'accès Suivi Cde'.\n"
    "- date_rdv : cherche dans cet ordre :\n"
    "  1. 'Rdv d'installation' ou 'RDV installation' (ignorer si marqué 'supprimé' ou 'annulé')\n"
    "  2. 'Date de livraison initiale', 'date de livraison', 'date d\'activation'\n"
    "  3. 'créneau', 'intervention prévue', 'date de pose'\n"
    "  Formats français à convertir en YYYY-MM-DDTHH:MM :\n"
    "  'le jeudi 25 juin' → déduire l'année depuis les autres dates du document → '2026-06-25T08:00'\n"
    "  '25/06' ou '25 juin' sans année → utiliser l'année visible ailleurs dans le document\n"
    "  Plage horaire '8h-12h' → prendre l'heure de début → T08:00\n"
    "  Si pas d'heure précise → T08:00 par défaut.\n"
    "- statut : analyser l'état de la commande et retourner EXACTEMENT l'une de ces valeurs :\n"
    "  'annule' si tu vois 'annulée', 'annulé', 'résiliée', 'résiliation', 'annulation'\n"
    "  'installe' si tu vois 'installée', 'installé', 'activée', 'activé', 'livrée', 'en service'\n"
    "  'confirme' si tu vois 'confirmée', 'confirmé', 'validée', 'en cours'\n"
    "  'no_show' si tu vois 'no show', 'absent', 'client absent'\n"
    "  'en_attente' dans tous les autres cas\n"
    "Retourne UNIQUEMENT le JSON brut, sans balises markdown."
)


async def _playwright_screenshot(url: str, login: str, password: str) -> bytes:
    from playwright.async_api import async_playwright
    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=True,
            args=["--no-sandbox", "--disable-dev-shm-usage", "--disable-gpu",
                  "--disable-setuid-sandbox", "--no-zygote"])
        ctx = await browser.new_context(viewport={"width": 1280, "height": 900},
            user_agent=("Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"))
        page = await ctx.new_page()
        try:
            await page.goto(url, wait_until="domcontentloaded", timeout=30000)
            await page.wait_for_timeout(2000)
            for _ in range(3):
                cur = page.url
                if not any(kw in cur for kw in ["login","auth","sso","signin","prelogin","portail","rso."]):
                    break
                for sel in ["#username","#login","#email",'input[name="username"]',
                            'input[name="login"]','input[type="email"]:visible','input[type="text"]:visible']:
                    try:
                        loc = page.locator(sel).first
                        if await loc.is_visible():
                            await loc.fill(login)
                            break
                    except Exception:
                        continue
                for sel in ["#password",'input[name="password"]','input[type="password"]:visible']:
                    try:
                        loc = page.locator(sel).first
                        if await loc.is_visible():
                            await loc.fill(password)
                            break
                    except Exception:
                        continue
                for sel in ['button[type="submit"]','input[type="submit"]',
                            "#bouton-valider",".btn-connexion",".btn-primary"]:
                    try:
                        loc = page.locator(sel).first
                        if await loc.is_visible():
                            await loc.click()
                            break
                    except Exception:
                        continue
                await page.wait_for_load_state("domcontentloaded", timeout=15000)
                await page.wait_for_timeout(2000)
            screenshot = await page.screenshot(full_page=False, type="jpeg", quality=88)
            return screenshot
        finally:
            await browser.close()


@app.route("/scan-affiche", methods=["POST"])
@login_required
def scan_affiche():
    import base64, json as _json
    photo = request.files.get("photo")
    if not photo:
        return jsonify({"error": "Aucune photo reçue"}), 400
    api_key = os.environ.get("ANTHROPIC_API_KEY")
    if not api_key:
        return jsonify({"error": "ANTHROPIC_API_KEY non configuré."}), 500
    img_bytes = photo.read()
    if len(img_bytes) > 10 * 1024 * 1024:
        return jsonify({"error": "Image trop lourde (max 10 Mo)."}), 400
    img_b64 = base64.standard_b64encode(img_bytes).decode("utf-8")
    media_type = photo.content_type if (photo.content_type or "").startswith("image/") else "image/jpeg"
    return _claude_vision(img_b64, media_type, api_key)


def _claude_vision(img_b64: str, media_type: str, api_key: str):
    import json as _json, anthropic
    try:
.        client = anthropic.Anthropic(api_key=api_key)
        msg = client.messages.create(
            model="claude-sonnet-4-6", max_tokens=1024,
            messages=[{"role": "user", "content": [
                {"type": "image", "source": {"type": "base64", "media_type": media_type, "data": img_b64}},
                {"type": "text", "text": _PROMPT_SCAN},
            ]}])
        txt = msg.content[0].text.strip()
        m = re.search(r'\{[\s\S]*\}', txt)
        txt = m.group() if m else txt
        return jsonify(_json.loads(txt))
    except _json.JSONDecodeError as e:
        return jsonify({"error": f"Impossible de lire la réponse IA : {e}"}), 500
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/track-commande", methods=["POST"])
@login_required
def track_commande():
    import base64
    url = request.form.get("url", "").strip()
    if not url:
        return jsonify({"error": "URL manquante."}), 400
    orange_login = os.environ.get("ORANGE_LOGIN", "")
    orange_password = os.environ.get("ORANGE_PASSWORD", "")
    if not orange_login or not orange_password:
        return jsonify({"error": "ORANGE_LOGIN et ORANGE_PASSWORD doivent être configurés dans les variables Railway."}), 500
    api_key = os.environ.get("ANTHROPIC_API_KEY")
    if not api_key:
        return jsonify({"error": "ANTHROPIC_API_KEY non configuré."}), 500
    try:
        loop = asyncio.new_event_loop()
        screenshot = loop.run_until_complete(_playwright_screenshot(url, orange_login, orange_password))
        loop.close()
        img_b64 = base64.standard_b64encode(screenshot).decode("utf-8")
        return _claude_vision(img_b64, "image/jpeg", api_key)
    except Exception as e:
        return jsonify({"error": f"Erreur Playwright : {e}"}), 500


@app.route("/offline")
def offline():
    return render_template("offline.html")


def _sort_numero(a):
    m = re.match(r'^(\d+)', str(a.numero).strip())
    return (int(m.group(1)) if m else 9999, str(a.numero))


def trouver_adresses_pour_rue(nom_session):
    nom_norm = nom_session.lower().strip()
    rues = [r[0] for r in db.session.query(AdresseImportee.rue).distinct().all()]
    matching = [r for r in rues if r.lower() in nom_norm or nom_norm in r.lower()]
    if not matching:
        return []
    adresses = AdresseImportee.query.filter(AdresseImportee.rue.in_(matching)).all()
    return sorted(adresses, key=_sort_numero)


@app.route("/prospection")
@login_required
def liste_prospection():
    sessions = (
        SessionProspection.query
        .order_by(SessionProspection.date.desc(), SessionProspection.created_at.desc()).all()
    )
    return render_template("prospection.html", sessions=sessions)


@app.route("/prospection/nouvelle", methods=["POST"])
@login_required
def nouvelle_session():
    nom = request.form.get("nom", "").strip()
    if not nom:
        flash("Donne un nom à ta session (ex: Rue Victor Hugo).", "warning")
        return redirect(url_for("liste_prospection"))
    sess = SessionProspection(nom=nom)
    db.session.add(sess)
    db.session.commit()
    return redirect(url_for("tap_session", session_id=sess.id))


@app.route("/prospection/<int:session_id>")
@login_required
def tap_session(session_id):
    sess = SessionProspection.query.get_or_404(session_id)
    adresses = trouver_adresses_pour_rue(sess.nom)
    if adresses:
        taps = {p.adresse_id: p.resultat for p in sess.portes if p.adresse_id is not None}
        adresses_data = [{
            "id": a.id, "rue": a.rue, "numero": a.numero,
            "complement": a.complement, "ville": a.ville or "",
            "resultat": taps.get(a.id),
        } for a in adresses]
        return render_template("tap_adresses.html", session=sess,
            adresses=adresses_data, resultats=RESULTATS_PORTE)
    return render_template("tap.html", session=sess, resultats=RESULTATS_PORTE)


@app.route("/prospection/<int:session_id>/itineraire")
@login_required
def itineraire_session(session_id):
    sess = SessionProspection.query.get_or_404(session_id)
    adresses = trouver_adresses_pour_rue(sess.nom)
    if not adresses:
        flash("Cette session n'a pas d'adresses importées.", "warning")
        return redirect(url_for("tap_session", session_id=session_id))
    taps = {p.adresse_id: p.resultat for p in sess.portes if p.adresse_id is not None}
    adresses_data = [{
        "id": a.id, "rue": a.rue, "numero": a.numero,
        "complement": a.complement or "", "ville": a.ville or "",
        "resultat": taps.get(a.id),
        "adresse_complete": f"{a.numero} {a.rue}{' ' + a.ville if a.ville else ''}",
    } for a in adresses]
    return render_template("itineraire.html", session=sess, adresses=adresses_data)


@app.route("/prospection/<int:session_id>/tap/<resultat>", methods=["POST"])
@login_required
def tap_porte(session_id, resultat):
    if resultat not in RESULTATS_PORTE:
        return jsonify({"error": "Résultat invalide"}), 400
    sess = SessionProspection.query.get_or_404(session_id)
    porte = Porte(session_id=session_id, resultat=resultat)
    db.session.add(porte)
    db.session.commit()
    return jsonify(sess.to_stats_dict())


@app.route("/prospection/<int:session_id>/annuler", methods=["POST"])
@login_required
def annuler_derniere_porte(session_id):
    sess = SessionProspection.query.get_or_404(session_id)
    derniere = Porte.query.filter_by(session_id=session_id).order_by(Porte.id.desc()).first()
    if derniere:
        db.session.delete(derniere)
        db.session.commit()
    return jsonify(sess.to_stats_dict())


@app.route("/prospection/<int:session_id>/tap-adresse/<int:adresse_id>/<resultat>", methods=["POST"])
@login_required
def tap_adresse_specifique(session_id, adresse_id, resultat):
    if resultat not in RESULTATS_PORTE:
        return jsonify({"error": "Résultat invalide"}), 400
    sess = SessionProspection.query.get_or_404(session_id)
    AdresseImportee.query.get_or_404(adresse_id)
    existing = Porte.query.filter_by(session_id=session_id, adresse_id=adresse_id).first()
    if existing:
        existing.resultat = resultat
    else:
        db.session.add(Porte(session_id=session_id, resultat=resultat, adresse_id=adresse_id))
    db.session.commit()
    all_taps = {p.adresse_id: p.resultat for p in sess.portes if p.adresse_id is not None}
    return jsonify({"stats": sess.to_stats_dict(), "adresse_id": adresse_id,
                    "resultat": resultat, "all_taps": all_taps})


@app.route("/prospection/<int:session_id>/supprimer", methods=["POST"])
@login_required
def supprimer_session(session_id):
    sess = SessionProspection.query.get_or_404(session_id)
    db.session.delete(sess)
    db.session.commit()
    flash("Session supprimée.", "info")
    return redirect(url_for("liste_prospection"))


@app.route("/import", methods=["GET", "POST"])
@login_required
def importer_fichier():
    if request.method == "POST":
        fichier = request.files.get("fichier")
        if not fichier or not fichier.filename:
            flash("Sélectionne un fichier Excel (.xlsx ou .xls).", "warning")
            return redirect(url_for("importer_fichier"))
        if not fichier.filename.lower().endswith((".xlsx", ".xls")):
            flash("Format invalide — seuls les fichiers .xlsx et .xls sont acceptés.", "warning")
            return redirect(url_for("importer_fichier"))
        try:
            from openpyxl import load_workbook
            import tempfile, os as _os
            tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
            fichier.save(tmp.name)
            tmp.close()
            wb = load_workbook(tmp.name, data_only=True)
            _os.unlink(tmp.name)
            ws = wb.worksheets[0]
            def cell_str(c):
                if c is None: return ""
                if hasattr(c, '__iter__') and not isinstance(c, str):
                    try: return "".join(getattr(p,'text',str(p)) for p in c).strip()
                    except Exception: pass
                if isinstance(c, bool): return ""
                if isinstance(c, float): return str(int(c)) if c.is_integer() else str(c)
                if isinstance(c, int): return str(c)
                return str(c).strip()
            all_rows = [
                [cell_str(c) for c in row]
                for row in ws.iter_rows(min_row=1, values_only=True)
                if any(c is not None for c in row)
            ]
            all_rows = [r for r in all_rows if any(v for v in r)]
            if not all_rows:
                flash("Le fichier est vide.", "warning")
                return redirect(url_for("importer_fichier"))
            n_cols = max(len(r) for r in all_rows)
            HEADER_KW = {"rue","adresse","voie","libelle","libellé","street",
                         "num","n°","no","porte","portes","numero","numéro","code",
                         "nom de voie","type voie","type de voie","libellé voie",
                         "nom","ville","profil","dep","immeuble","fermeture"}
            first_low = all_rows[0][0].lower().strip() if all_rows[0] else ""
            has_header = first_low in HEADER_KW
            header_row = all_rows[0] if has_header else None
            data_rows  = all_rows[1:] if has_header else all_rows
            if not data_rows:
                flash("Aucune donnée trouvée dans le fichier.", "warning")
                return redirect(url_for("importer_fichier"))
            PREFIXES_RUE = (
                "rue ","avenue ","av ","av. ","boulevard ","bd ","bd.",
                "chemin ","impasse ","allée ","allee ","passage ","place ",
                "route ","voie ","résidence ","residence ","cité ","cite ",
                "square ","villa ","domaine ","lot ","lieu-dit",
            )
            sample = data_rows[:50]
            scores_rue = [0]*n_cols; scores_num = [0]*n_cols; scores_ville = [0]*n_cols
            col_unique_vals = [set() for _ in range(n_cols)]
            col_zero_count  = [0]*n_cols
            for row in sample:
                for j in range(min(n_cols,len(row))):
                    v = row[j]
                    if not v: continue
                    vl = v.lower()
                    if any(vl.startswith(p) for p in PREFIXES_RUE): scores_rue[j] += 4
                    elif " " in v and not v[0].isdigit() and "/" not in v and len(v)>5: scores_rue[j] += 1
                    if "/" not in v and re.match(r'^\d{1,4}\w{0,3}$',v): scores_num[j] += 4
                    elif re.match(r'^\d+$',v) and len(v)<=4: scores_num[j] += 2
                    if re.match(r'^[A-Za-zÀ-ÿ\s\-\']+$',v) and 2<=len(v)<=35 and "/" not in v:
                        if not any(vl.startswith(p) for p in PREFIXES_RUE): scores_ville[j] += 2
                    if v=="0": col_zero_count[j] += 1
                    else: col_unique_vals[j].add(v)
            for j in range(n_cols):
                n_unique=len(col_unique_vals[j]); n_zero=col_zero_count[j]; total=n_unique+n_zero
                if total==0: continue
                if total>3 and n_zero/total>0.5: scores_num[j]=0
                if n_unique>=3: scores_num[j]+=n_unique*2
            col_rue=col_num=col_ville=None
            if has_header and header_row:
                for j,c in enumerate(header_row):
                    cl=c.lower()
                    if col_ville is None and any(w in cl for w in ("ville","commune","localit","city")): col_ville=j
                    if col_rue is None and "adresse" in cl and "mail" not in cl and "email" not in cl: col_rue=j
                    if col_num is None and re.search(r'num[eé]ro\s*rue|n°\s*rue|num\s*rue',cl): col_num=j
                for j,c in enumerate(header_row):
                    cl=c.lower()
                    if col_rue is None and any(w in cl for w in ("rue","voie","libelle","libellé")): col_rue=j
                    if col_num is None and any(w in cl for w in ("num","n°","numéro","numero")): col_num=j
                for j,c in enumerate(header_row):
                    cl=c.lower()
                    if col_rue is None and "nom" in cl: col_rue=j
            if col_rue is None: col_rue=max(range(n_cols),key=lambda j:scores_rue[j])
            if col_num is None:
                candidates=[j for j in range(n_cols) if j!=col_rue]
                col_num=max(candidates,key=lambda j:scores_num[j]) if candidates else (1 if col_rue==0 else 0)
            if col_ville is None:
                candidates_v=[j for j in range(n_cols) if j not in (col_rue,col_num)]
                if candidates_v:
                    best_v=max(candidates_v,key=lambda j:scores_ville[j])
                    if scores_ville[best_v]>=4: col_ville=best_v
            nb_ok=nb_skip=0
            for cells in data_rows:
                rue_val=cells[col_rue] if col_rue<len(cells) else ""
                num_val=cells[col_num] if col_num<len(cells) else ""
                if not num_val and rue_val:
                    m=re.match(r'^(\d+\w*)\s+(.+)$',rue_val)
                    if m: num_val,rue_val=m.group(1),m.group(2)
                if not rue_val or not num_val:
                    nb_skip+=1; continue
                ville_val=cells[col_ville] if col_ville is not None and col_ville<len(cells) else None
                ville_val=ville_val or None
                db.session.add(AdresseImportee(rue=rue_val,numero=num_val,ville=ville_val))
                nb_ok+=1
            db.session.commit()
            col_info=f"col {col_rue+1}=Rue, col {col_num+1}=Numéro"
            if col_ville is not None: col_info+=f", col {col_ville+1}=Ville"
            flash(f"{nb_ok} adresses importées ({col_info}). {nb_skip} lignes ignorées.","success")
            return redirect(url_for("importer_fichier"))
        except Exception as exc:
            db.session.rollback()
            flash(f"Erreur lors de l'import : {exc}", "danger")
    nb_adresses = AdresseImportee.query.count()
    from sqlalchemy import func as sa_func
    rues = (
        db.session.query(AdresseImportee.rue, AdresseImportee.ville,
            sa_func.count(AdresseImportee.id).label("nb"))
        .group_by(AdresseImportee.rue, AdresseImportee.ville)
        .order_by(AdresseImportee.ville, AdresseImportee.rue).all()
    )
    villes = sorted(set(r.ville for r in rues if r.ville))
    return render_template("import.html", nb_adresses=nb_adresses,
        nb_rues=len(rues), nb_villes=len(villes), rues=rues)


@app.route("/import/structure", methods=["POST"])
@login_required
def structure_fichier():
    fichier = request.files.get("fichier")
    if not fichier: return "Aucun fichier", 400
    try:
        from openpyxl import load_workbook
        import tempfile, os as _os
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
        fichier.save(tmp.name); tmp.close()
        wb = load_workbook(tmp.name, data_only=True); _os.unlink(tmp.name)
        ws = wb.worksheets[0]
        html = """<html><head><meta charset="utf-8"><style>body{font-family:monospace;font-size:13px;padding:20px}
        table{border-collapse:collapse}td,th{border:1px solid #ccc;padding:4px 8px;max-width:200px;overflow:hidden;text-overflow:ellipsis;white-space:nowrap}
        th{background:#eee}.row-num{color:#999;font-size:11px}.type{color:#07c;font-size:10px}</style></head><body>
        <h3>Structure du fichier (15 premières lignes)</h3>
        <p>Chaque cellule montre : <strong>valeur</strong> <span style="color:#07c">type Python</span></p>
        <table><tr><th>#</th>"""
        first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), [])
        for j in range(len(first_row)): html += f"<th>Col {j+1}</th>"
        html += "</tr>"
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True)):
            html += f"<tr><td class='row-num'>{i+1}</td>"
            for c in row:
                typ = type(c).__name__
                val = str(c)[:60] if c is not None else "(vide)"
                html += f"<td>{val}<br><span class='type'>{typ}</span></td>"
            html += "</tr>"
        html += "</table></body></html>"
        return html
    except Exception as e:
        return f"Erreur : {e}", 500


@app.route("/repasser")
@login_required
def repasser():
    from collections import defaultdict
    rows_addr = (
        db.session.query(Porte, AdresseImportee, SessionProspection)
        .join(AdresseImportee, Porte.adresse_id == AdresseImportee.id)
        .join(SessionProspection, Porte.session_id == SessionProspection.id)
        .filter(Porte.resultat == "absent")
        .order_by(SessionProspection.date.desc(), AdresseImportee.rue, AdresseImportee.numero).all()
    )
    rows_libre = (
        db.session.query(Porte, SessionProspection)
        .join(SessionProspection, Porte.session_id == SessionProspection.id)
        .filter(Porte.resultat == "absent", Porte.adresse_id == None)
        .order_by(SessionProspection.date.desc()).all()
    )
    JOURS = ["Lundi","Mardi","Mercredi","Jeudi","Vendredi","Samedi","Dimanche"]
    MOIS = ["","jan.","fév.","mars","avr.","mai","juin","juil.","août","sep.","oct.","nov.","déc."]
    grouped = defaultdict(lambda: {"adresses": [], "libres": defaultdict(int)})
    for porte, adresse, sess in rows_addr:
        grouped[sess.date]["adresses"].append((porte, adresse, sess))
    for porte, sess in rows_libre:
        grouped[sess.date]["libres"][sess.nom] += 1
    groupes = []
    for d in sorted(grouped.keys(), reverse=True):
        label = f"{JOURS[d.weekday()]} {d.day} {MOIS[d.month]} {d.year}"
        groupes.append((label, grouped[d]["adresses"], dict(grouped[d]["libres"])))
    nb_total = len(rows_addr) + len(rows_libre)
    return render_template("repasser.html", groupes=groupes, nb_total=nb_total)


@app.route("/rappels")
@login_required
def liste_rappels():
    rappels = Rappel.query.filter_by(done=False).order_by(Rappel.created_at).all()
    faits = Rappel.query.filter_by(done=True).order_by(Rappel.created_at.desc()).limit(10).all()
    return render_template("rappels.html", rappels=rappels, faits=faits, moments=MOMENTS_RAPPEL)


@app.route("/rappels/nouveau", methods=["POST"])
@login_required
def nouveau_rappel():
    nom = request.form.get("nom", "").strip()
    telephone = request.form.get("telephone", "").strip() or None
    motif = request.form.get("motif", "").strip()
    moment = request.form.get("moment", "pause")
    if not nom or not motif:
        flash("Nom et motif obligatoires.", "warning")
        return redirect(url_for("liste_rappels"))
    db.session.add(Rappel(nom=nom, telephone=telephone, motif=motif, moment=moment))
    db.session.commit()
    flash(f"Rappel ajouté pour {nom}.", "success")
    return redirect(url_for("liste_rappels"))


@app.route("/rappels/<int:rappel_id>/done", methods=["POST"])
@login_required
def rappel_done(rappel_id):
    r = Rappel.query.get_or_404(rappel_id)
    r.done = True
    db.session.commit()
    return redirect(url_for("liste_rappels"))


@app.route("/rappels/<int:rappel_id>/supprimer", methods=["POST"])
@login_required
def rappel_supprimer(rappel_id):
    r = Rappel.query.get_or_404(rappel_id)
    db.session.delete(r)
    db.session.commit()
    return redirect(url_for("liste_rappels"))


@app.route("/import/effacer", methods=["POST"])
@login_required
def effacer_adresses():
    AdresseImportee.query.delete()
    db.session.commit()
    flash("Toutes les adresses importées ont été supprimées.", "info")
    return redirect(url_for("importer_fichier"))


@app.route("/backup-db")
@login_required
def backup_db():
    import re as _re
    db_uri = app.config["SQLALCHEMY_DATABASE_URI"]
    m = _re.match(r"sqlite:///+(.*)", db_uri)
    if not m:
        flash("Backup non disponible (base non SQLite).", "danger")
        return redirect(url_for("dashboard"))
    db_file = "/" + m.group(1).lstrip("/")
    if not os.path.exists(db_file):
        flash(f"Fichier base introuvable : {db_file}", "danger")
        return redirect(url_for("dashboard"))
    return send_file(db_file, as_attachment=True, download_name="ventes_backup.db",
        mimetype="application/octet-stream")


@app.route("/restore-db", methods=["GET", "POST"])
@login_required
def restore_db():
    import re as _re, shutil as _shutil, tempfile as _tempfile
    if request.method == "POST":
        f = request.files.get("fichier_db")
        if not f or not f.filename.endswith(".db"):
            flash("Fichier .db requis.", "danger")
            return redirect(url_for("restore_db"))
        db_uri = app.config["SQLALCHEMY_DATABASE_URI"]
        m = _re.match(r"sqlite:///+(.*)", db_uri)
        if not m:
            flash("Restore non disponible (base non SQLite).", "danger")
            return redirect(url_for("dashboard"))
        db_file = "/" + m.group(1).lstrip("/")
        if os.path.exists(db_file): _shutil.copy2(db_file, db_file + ".bak")
        tmp = _tempfile.NamedTemporaryFile(delete=False, suffix=".db")
        f.save(tmp.name); tmp.close()
        _shutil.move(tmp.name, db_file)
        flash("Base restaurée. Redémarre le service Railway pour recharger les données.", "success")
        return redirect(url_for("dashboard"))
    return render_template("restore_db.html")


@app.route("/memoire")
@login_required
def memoire():
    from sqlalchemy import func
    sessions = SessionProspection.query.order_by(SessionProspection.date.desc()).all()
    ventes = Vente.query.order_by(Vente.date_signature.desc()).all()
    timeline = {}
    for s in sessions:
        d = s.date
        if d not in timeline: timeline[d] = {"portes": 0, "ventes": [], "zones": []}
        timeline[d]["portes"] += s.total
        if s.nom not in timeline[d]["zones"]: timeline[d]["zones"].append(s.nom)
    for v in ventes:
        d = v.date_signature
        if d not in timeline: timeline[d] = {"portes": 0, "ventes": [], "zones": []}
        timeline[d]["ventes"].append(v)
    timeline_sorted = sorted(timeline.items(), key=lambda x: x[0], reverse=True)
    villes = db.session.query(
        AdresseImportee.ville, func.count(AdresseImportee.id).label("nb"),
    ).filter(AdresseImportee.ville != None, AdresseImportee.ville != "").group_by(
        AdresseImportee.ville
    ).order_by(func.count(AdresseImportee.id).desc()).all()
    adresses_contrats = db.session.query(
        Vente.adresse, func.count(Vente.id).label("nb"),
    ).filter(Vente.adresse != None, Vente.adresse != "").group_by(
        Vente.adresse
    ).order_by(func.count(Vente.id).desc()).all()
    return render_template("memoire.html", timeline=timeline_sorted,
        villes=villes, adresses_contrats=adresses_contrats)


# ---------------------------------------------------------------------------
# Route temporaire : Seed données février 2026
# ---------------------------------------------------------------------------

@app.route("/admin/seed-fevrier")
@login_required
def seed_fevrier():
    from datetime import date as _d, datetime as _dt, timedelta as _td
    import random as _rnd
    _rnd.seed(2026)
    if Vente.query.filter(
        Vente.date_signature >= _d(2026, 2, 1),
        Vente.date_signature <= _d(2026, 2, 28),
    ).count() > 0:
        flash("Donnees fevrier deja injectees.", "warning")
        return redirect(url_for("dashboard"))
    PRENOMS = ["Mohamed", "Jean", "Pierre", "Ahmed", "Thomas", "Nicolas", "Karim", "David",
               "Laurent", "Eric", "Rachid", "Patrick", "Stephane", "Bruno", "Marc", "Francois",
               "Didier", "Sebastien", "Julien", "Anthony", "Christophe", "Alexandre", "Guillaume",
               "Romain", "Maxime", "Sophie", "Marie", "Fatima", "Isabelle", "Nathalie", "Celine",
               "Sandrine", "Amina", "Sylvie", "Virginie", "Aurelie", "Laure", "Karine", "Valerie",
               "Nadege", "Leila", "Zineb", "Hafida", "Samia", "Karima", "Yasmine", "Sonia", "Nadia", "Aicha"]
    NOMS = ["MARTIN", "BERNARD", "DUBOIS", "THOMAS", "ROBERT", "RICHARD", "PETIT", "DURAND",
            "MOREAU", "SIMON", "LAURENT", "LEFEBVRE", "MICHEL", "GARCIA", "DAVID", "BERTRAND",
            "ROUX", "VINCENT", "FOURNIER", "MOREL", "GIRARD", "ANDRE", "LEROY", "DUPONT",
            "LAMBERT", "BONNET", "FRANCOIS", "MARTINEZ", "LEGRAND", "GARNIER", "FAURE", "ROUSSEAU",
            "BLANC", "GUERIN", "MULLER", "HENRY", "ROUSSEL", "NICOLAS", "PERRIN", "MORIN",
            "MATHIEU", "CLEMENT", "GAUTHIER", "DUMONT", "LOPEZ", "FONTAINE", "CHEVALIER", "ROBIN",
            "MASSON"]
    ADRESSES = [
        "12 Rue Victor Hugo, Nimes", "45 Avenue Jean Jaures, Nimes", "8 Boulevard Gambetta, Nimes",
        "23 Rue de la Paix, Nimes", "67 Allee des Roses, Nimes", "3 Impasse du Moulin, Nimes",
        "15 Rue du Commerce, Nimes", "89 Avenue de la Republique, Nimes", "5 Chemin des Lilas, Nimes",
        "34 Rue Saint-Nicolas, Nimes", "71 Boulevard du Marechal Foch, Ales", "18 Rue des Ecoles, Ales",
        "56 Avenue Pierre Mendes France, Nimes", "9 Rue Emile Zola, Nimes", "42 Impasse des Acacias, Nimes",
        "27 Rue du General de Gaulle, Uzes", "63 Avenue Francois Mitterrand, Nimes", "11 Rue Pasteur, Nimes",
        "48 Boulevard Louis Blanc, Nimes", "7 Allee des Chenes, Nimes", "33 Rue Lamartine, Nimes",
        "80 Avenue de la Gare, Ales", "16 Rue Voltaire, Nimes", "55 Boulevard Raspail, Nimes",
        "22 Rue du 8 Mai 1945, Nimes", "41 Avenue de la Liberation, Lunel", "6 Impasse des Muriers, Nimes",
        "77 Rue Jules Ferry, Nimes", "14 Chemin du Moulin, Nimes", "39 Rue Aristide Briand, Nimes",
        "25 Rue de la Fontaine, Montpellier", "58 Avenue du Pont, Nimes", "13 Rue des Fleurs, Nimes",
        "47 Allee des Pins, Nimes", "2 Rue de la Mairie, Vergeze", "31 Avenue des Sports, Nimes",
        "64 Rue de la Croix, Nimes", "19 Impasse des Vignes, Nimes", "86 Boulevard de la Paix, Nimes",
        "10 Rue Moliere, Nimes", "43 Avenue Carnot, Nimes", "28 Rue de l'Eglise, Saint-Gilles",
        "72 Allee des Platanes, Nimes", "4 Chemin de la Colline, Nimes", "37 Rue Gambetta, Lunel",
        "61 Boulevard National, Nimes", "24 Rue de la Liberte, Nimes", "50 Avenue du Midi, Nimes",
        "17 Rue des Amandiers, Nimes",
    ]
    PRODUITS = (["Livebox Fibre"] * 29 + ["Livebox Up"] * 10 + ["Livebox Max"] * 7 + ["Serie Special Lite Fibre"] * 3)
    DATA = [
        ("2026-02-02", ["installe", "confirme", "confirme"]),
        ("2026-02-03", ["installe", "confirme", "no_show"]),
        ("2026-02-04", ["installe", "confirme"]),
        ("2026-02-05", ["confirme", "confirme"]),
        ("2026-02-06", ["confirme"]),
        ("2026-02-09", ["installe", "installe", "confirme"]),
        ("2026-02-10", ["installe", "installe", "annule"]),
        ("2026-02-11", ["installe", "confirme", "confirme"]),
        ("2026-02-12", ["installe", "confirme", "no_show"]),
        ("2026-02-13", ["confirme", "confirme"]),
        ("2026-02-16", ["installe", "installe", "confirme"]),
        ("2026-02-17", ["installe", "installe", "confirme"]),
        ("2026-02-18", ["installe", "installe", "no_show"]),
        ("2026-02-19", ["installe", "confirme", "confirme"]),
        ("2026-02-20", ["installe", "confirme", "no_show"]),
        ("2026-02-23", ["installe", "confirme"]),
        ("2026-02-24", ["installe", "confirme"]),
        ("2026-02-25", ["confirme", "confirme"]),
        ("2026-02-26", ["confirme", "no_show"]),
        ("2026-02-27", ["confirme"]),
    ]
    heures = [8, 9, 10, 14, 15, 16]
    idx = 0
    for date_str, statuts in DATA:
        sig = _d.fromisoformat(date_str)
        for statut in statuts:
            rdv = _dt.combine(sig + _td(days=_rnd.randint(4, 12)),
                              _dt.min.time().replace(hour=heures[idx % len(heures)]))
            db.session.add(Vente(
                prenom=PRENOMS[idx % len(PRENOMS)],
                nom=NOMS[idx % len(NOMS)],
                telephone=f"+336{50000000 + idx:08d}",
                adresse=ADRESSES[idx],
                produit=PRODUITS[idx],
                date_rdv=rdv,
                date_signature=sig,
                statut=statut,
                sms_envoye=(statut not in ("en_attente", "confirme")),
            ))
            idx += 1
    db.session.commit()
    flash("49 ventes de fevrier 2026 injectees avec succes !", "success")
    return redirect(url_for("dashboard"))


@app.route("/admin/seed-mars")
@login_required
def seed_mars():
    from datetime import date as _d, datetime as _dt, timedelta as _td
    import random as _rnd
    _rnd.seed(2027)
    if Vente.query.filter(
        Vente.date_signature >= _d(2026, 3, 1),
        Vente.date_signature <= _d(2026, 3, 31),
    ).count() > 0:
        flash("Donnees mars deja injectees.", "warning")
        return redirect(url_for("dashboard"))
    PRENOMS = ["Karim", "Sophie", "Laurent", "Nathalie", "Mohamed", "Isabelle", "David",
               "Virginie", "Nicolas", "Fatima", "Sebastien", "Celine", "Thomas", "Amina",
               "Guillaume", "Sylvie", "Romain", "Zineb", "Maxime", "Karine", "Pierre",
               "Sandrine", "Eric", "Leila", "Julien", "Valerie", "Alexandre"]
    NOMS = ["MARTIN", "DUBOIS", "THOMAS", "ROBERT", "RICHARD", "PETIT", "DURAND",
            "MOREAU", "SIMON", "LAURENT", "LEFEBVRE", "MICHEL", "GARCIA", "DAVID",
            "BERTRAND", "ROUX", "VINCENT", "FOURNIER", "MOREL", "GIRARD", "ANDRE",
            "LEROY", "DUPONT", "LAMBERT", "BONNET", "FRANCOIS", "MARTINEZ"]
    ADRESSES = [
        "3 Rue du Gard, Nimes", "18 Avenue des Arenes, Nimes", "27 Rue Plechat, Nimes",
        "52 Boulevard Victor Hugo, Nimes", "9 Impasse des Oliviers, Nimes",
        "34 Rue de la Fontaine, Nimes", "71 Avenue du Marche, Ales",
        "12 Rue Colbert, Nimes", "45 Chemin des Pins, Nimes", "6 Rue du Temple, Uzes",
        "83 Boulevard Sergent Triaire, Nimes", "21 Rue de la Baume, Nimes",
        "58 Avenue du Grau, Lunel", "15 Rue Bigot, Nimes", "39 Impasse des Cedres, Nimes",
        "66 Rue des Tuileries, Nimes", "8 Avenue de Montpellier, Nimes",
        "43 Rue de la Croix de Fer, Nimes", "29 Boulevard Jean Jaures, Ales",
        "74 Rue des Asphodeles, Nimes", "11 Chemin de la Gardiole, Vergeze",
        "55 Avenue du President Wilson, Nimes", "22 Rue Porte de France, Nimes",
        "37 Impasse des Amandiers, Nimes", "90 Rue Nationale, Lunel",
        "4 Rue des Jacobins, Nimes", "68 Boulevard de la Liberation, Nimes",
    ]
    PRODUITS = (["Livebox Fibre"] * 16 + ["Livebox Up"] * 7 + ["Livebox Max"] * 3 + ["Série Spécial Lite Fibre"] * 1)
    DATA = [
        ("2026-03-02", ["installe", "confirme", "confirme"]),
        ("2026-03-03", ["installe", "confirme", "confirme"]),
        ("2026-03-04", ["installe", "confirme"]),
        ("2026-03-05", ["annule", "confirme", "confirme"]),
        ("2026-03-06", ["confirme", "confirme"]),
        ("2026-03-09", ["installe", "confirme", "confirme"]),
        ("2026-03-10", ["installe", "confirme", "confirme"]),
        ("2026-03-11", ["installe", "confirme", "confirme"]),
        ("2026-03-12", ["confirme", "confirme", "confirme"]),
        ("2026-03-13", ["confirme", "confirme"]),
    ]
    heures = [8, 9, 10, 14, 15, 16]
    idx = 0
    for date_str, statuts in DATA:
        sig = _d.fromisoformat(date_str)
        for statut in statuts:
            rdv = _dt.combine(sig + _td(days=_rnd.randint(4, 14)),
                              _dt.min.time().replace(hour=heures[idx % len(heures)]))
            db.session.add(Vente(
                prenom=PRENOMS[idx % len(PRENOMS)],
                nom=NOMS[idx % len(NOMS)],
                telephone=f"+336{60000000 + idx:08d}",
                adresse=ADRESSES[idx % len(ADRESSES)],
                produit=PRODUITS[idx % len(PRODUITS)],
                date_rdv=rdv,
                date_signature=sig,
                statut=statut,
                sms_envoye=(statut not in ("en_attente", "confirme")),
            ))
            idx += 1
    db.session.commit()
    flash("27 ventes de mars 2026 injectees avec succes !", "success")
    return redirect(url_for("dashboard"))


@app.route("/admin/salaire", methods=["POST"])
@login_required
def sauvegarder_salaire():
    try:
        annee = int(request.form["annee"])
        mois = int(request.form["mois"])
        montant = float(request.form["montant"].replace(",", ".").replace(" ", "").replace(" ", ""))
        existing = SalaireMensuel.query.filter_by(annee=annee, mois=mois).first()
        if existing:
            existing.montant_net = montant
        else:
            db.session.add(SalaireMensuel(annee=annee, mois=mois, montant_net=montant))
        db.session.commit()
        flash("Salaire enregistré.", "success")
    except Exception as e:
        flash(f"Erreur : {e}", "danger")
    return redirect(url_for("stats"))


# ---------------------------------------------------------------------------
# Démarrage
# ---------------------------------------------------------------------------

def creer_scheduler():
    scheduler = BackgroundScheduler()
    scheduler.add_job(envoyer_rappels_du_jour, trigger="cron", hour=9, minute=0,
        id="rappels_sms", replace_existing=True)
    scheduler.add_job(envoyer_sms_rappels_clients, trigger="cron", hour=12, minute=30,
        id="sms_rappels_pause", replace_existing=True)
    scheduler.add_job(envoyer_sms_rappels_clients, trigger="cron", hour=19, minute=0,
        id="sms_rappels_soir", replace_existing=True)
    scheduler.start()
    return scheduler


with app.app_context():
    db.create_all()
    from sqlalchemy import text, inspect as sa_inspect
    _insp = sa_inspect(db.engine)
    _cols = [c["name"] for c in _insp.get_columns("portes")]
    if "adresse_id" not in _cols:
        with db.engine.connect() as _conn:
            _conn.execute(text("ALTER TABLE portes ADD COLUMN adresse_id INTEGER"))
            _conn.commit()
    _cols_ai = [c["name"] for c in _insp.get_columns("adresses_importees")]
    if "ville" not in _cols_ai:
        with db.engine.connect() as _conn:
            _conn.execute(text("ALTER TABLE adresses_importees ADD COLUMN ville VARCHAR(100)"))
            _conn.commit()

scheduler = creer_scheduler()

if __name__ == "__main__":
    try:
        app.run(debug=False, host="0.0.0.0", port=5000)
    finally:
        scheduler.shutdown()
